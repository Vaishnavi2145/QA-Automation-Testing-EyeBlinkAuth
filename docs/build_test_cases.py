import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Test Cases"

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

headers = [
    "Test Case ID", "Module", "Type", "Scenario", "Precondition",
    "Test Steps", "Test Data", "Expected Result", "Priority", "Severity",
    "Automation Status", "Notes"
]

rows = [
["TC-SU-01","Signup","Positive","Successful signup with valid details and matching blink count",
 "App running; MySQL up; user email not already registered",
 "1. Go to /signup\n2. Fill Name, Email, Phone, Gender, Address\n3. Click Start Camera, blink N times\n4. Click Submit & Save",
 "Name=Test User, Email=unique@example.com, blink=3",
 "Face samples captured (8), row inserted in users table, redirected to /login with 'Signup successful!' message",
 "P1","N/A (positive)","Manual","Cannot be automated - see BUG-01 (server-side blocking camera capture)"],

["TC-SU-02","Signup","Negative","Submit with Name blank",
 "On /signup page","1. Leave Name blank\n2. Fill other required fields\n3. Submit",
 "Name=''","'Please complete the form and capture your blink count.' flash, stays on /signup",
 "P2","Medium","Automated (API)","api.SignupApiTests#missingNameRejected"],

["TC-SU-03","Signup","Negative","Submit with Email blank",
 "On /signup page","1. Leave Email blank\n2. Fill other fields\n3. Submit",
 "Email=''","Same 'please complete the form' error","P2","Medium",
 "Automated (API)","api.SignupApiTests#missingEmailRejected"],

["TC-SU-04","Signup","Negative","Submit with non-numeric blink_count",
 "On /signup page (simulated - hidden field manipulated)","1. Set blink_count to a non-numeric value\n2. Submit",
 "blink_count='abc'","Rejected - blink_count.isdigit() is False","P2","Medium",
 "Automated (API)","api.SignupApiTests#nonNumericBlinkCountRejected"],

["TC-SU-05","Signup","Boundary","Submit with blink_count = '-1'",
 "On /signup page","1. Set blink_count to '-1'\n2. Submit","blink_count='-1'",
 "Rejected - Python's '-1'.isdigit() is False (minus sign isn't a digit)","P3","Low",
 "Automated (API)","api.SignupApiTests#negativeBlinkCountRejected"],

["TC-SU-06","Signup","Negative","Submit with duplicate email",
 "A user with this email already exists in the users table",
 "1. Fill form with an already-registered email\n2. Complete camera capture\n3. Submit",
 "Email = an existing user's email","Expected: 'Email already exists. Try logging in.' " 
 "ACTUAL (see BUG-04): db/blinkauth.sql defines no UNIQUE constraint on users.email, "
 "so the insert likely succeeds silently instead of raising IntegrityError",
 "P1","High","Manual","Needs live DB check - see DEFECT_LOG BUG-04"],

["TC-SU-07","Signup","Negative","Multiple faces in camera frame during capture",
 "On /signup, two people in front of the camera","1. Start camera with 2+ faces visible\n2. Attempt capture",
 "N/A","UI shows '❌ Multiple faces detected'; capture does not proceed until exactly one face is visible",
 "P3","Low","Manual","Requires physical camera / multiple people"],

["TC-SU-08","Signup","Negative","No face detected during capture",
 "On /signup, camera pointed away from any face","1. Start camera with no face in frame",
 "N/A","UI shows 'No face detected'; capture does not proceed","P3","Low","Manual",
 "Requires physical camera"],

["TC-LG-01","Login","Positive","Successful login with correct email, face, and blink count",
 "User previously completed signup","1. Go to /login\n2. Enter registered email\n3. Start camera, blink correct count\n4. Submit",
 "Email + correct blink count + live matching face","Redirected to /dashboard, 'Login successful! Face + Blink verified.'",
 "P1","N/A (positive)","Manual","Cannot be automated - needs a real registered face, see LoginPage javadoc"],

["TC-LG-02","Login","Negative","Login with unregistered email",
 "Email does not exist in users table","1. Enter an unregistered email\n2. Provide any blink/face data\n3. Submit",
 "Email='no.such.user@example.com'","'Email not registered.' flash, stays on /login","P1","Medium",
 "Automated (API)","api.LoginApiTests#unregisteredEmailRejected"],

["TC-LG-03","Login","Negative","Login with blank email",
 "On /login page","1. Leave email blank\n2. Submit","Email=''","'Incomplete login data.' flash",
 "P2","Medium","Automated (API)","api.LoginApiTests#missingEmailRejected"],

["TC-LG-04","Login","Negative","Login with blank blink_count",
 "On /login page","1. Leave blink_count blank (skip camera)\n2. Submit","blink_count=''",
 "'Incomplete login data.' flash","P2","Medium","Automated (API)","api.LoginApiTests#missingBlinkCountRejected"],

["TC-LG-05","Login","Negative","Click Login before starting camera",
 "On /login page, camera not started","1. Enter email\n2. Click Login without clicking Start Camera",
 "N/A","Client-side alert 'Camera not ready. Please start camera.'; form does not submit",
 "P2","Low","Automated (UI)","tests.LoginUiTests#loginBlockedWithoutCamera"],

["TC-LG-06","Login","Negative","Face does not match registered face",
 "User registered; different person / photo attempts login","1. Enter registered user's email\n2. Use a different face in front of camera\n3. Blink any count\n4. Submit",
 "Mismatched face","'Face does not match the entered email.' flash; notification email attempt fired (see BUG-05)",
 "P1","High","Manual","Needs a second real face sample to test against"],

["TC-LG-07","Login","Negative","Correct face, wrong blink count",
 "User registered with known blink_count","1. Enter registered email\n2. Use correct face\n3. Blink a different number of times\n4. Submit",
 "blink_count != stored value","'Blink count does not match.' flash; notification email attempt fired",
 "P1","High","Manual","Needs a real registered face; proves 2FA layer is enforced"],

["TC-LG-08","Login","Security","Blink count is trusted from a client-submitted hidden field",
 "User registered","1. Log in normally up to the camera step\n2. Using browser dev tools, edit the hidden #blinkInput value to the CORRECT stored count without actually blinking\n3. Submit with a valid face capture",
 "Manipulated blink_count matching DB value, no real blinks performed",
 "Login likely SUCCEEDS - the server only compares the submitted number to the stored number, it never verifies blinks actually happened client-side. See BUG-09 (business-logic weakness): the blink factor can be spoofed by anyone who can edit page state, defeating the second factor entirely.",
 "P1","High","Manual","Requires browser dev tools; documents a business-logic (not just UI) weakness"],

["TC-SEC-01","Security","Security","Signup form has no CSRF token",
 "On /signup page","1. Inspect form for hidden CSRF token field","N/A",
 "No CSRF token present - forms are vulnerable to Cross-Site Request Forgery",
 "P2","Medium","Automated (UI)","tests.SignupFormValidationTests#signupFormHasNoCsrfToken"],

["TC-SEC-02","Security","Security","Login form has no CSRF token",
 "On /login page","1. Inspect form for hidden CSRF token field","N/A","Same as TC-SEC-01","P2","Medium",
 "Automated (UI)","tests.LoginUiTests#loginFormHasNoCsrfToken"],

["TC-SEC-03","Security","Security","/dashboard requires an active session",
 "No prior login (fresh browser / no cookies)","1. Navigate directly to /dashboard",
 "N/A","Redirected to /login with 'Please login first.'","P1","N/A (control - passes)",
 "Automated (UI + API)","tests.AccessControlUiTests#dashboardRequiresLogin, api.AccessControlApiTests#dashboardRequiresSession"],

["TC-SEC-04","Security","Security","/files requires an active session",
 "No prior login","1. Navigate directly to /files","N/A","Redirected to /login","P1","N/A (control - passes)",
 "Automated (UI + API)","tests.AccessControlUiTests#filesRequiresLogin, api.AccessControlApiTests#filesRequiresSession"],

["TC-SEC-05","Security","Security","/upload requires an active session",
 "No prior login","1. POST directly to /upload with a file, no session cookie","N/A",
 "Redirected to /login","P1","N/A (control - passes)","Automated (API)","api.AccessControlApiTests#uploadRequiresSession"],

["TC-SEC-06","Security","Security (Critical)","/download/<email>/<filename> has NO session check - IDOR",
 "Any file has ever been uploaded by any user","1. With no login/session at all, request GET /download/<victim_email>/<real_filename> directly",
 "A known email + filename combination from static/uploads/",
 "BUG: file is returned with HTTP 200 and no authentication - any user's files can be downloaded by anyone who knows/guesses their email and filenames",
 "P1","Critical","Automated (API)","api.AccessControlApiTests#downloadWithoutSessionSucceeds - see DEFECT_LOG BUG-03"],

["TC-SEC-07","Security","Security","face_image with no comma crashes the login request",
 "Registered user exists","1. POST to /login with face_image='not-a-data-url' (no comma), valid email + blink_count",
 "face_image without a comma","Expected: graceful error. Actual (per code review): unhandled ValueError -> HTTP 500",
 "P2","Medium","Automated (API)","api.LoginApiTests#malformedFaceImageCausesServerError - see DEFECT_LOG BUG-08"],

["TC-SEC-08","Security","Security","Login crashes if no user has ever signed up (ref_embed.pkl missing)",
 "Fresh install / ref_embed.pkl deleted or never created","1. Attempt login with any registered email + any blink/face data",
 "N/A","Expected: 'No face registered for this email.' Actual (per code review): unhandled FileNotFoundError -> HTTP 500 " 
 "because open(REF_EMBED_PATH) isn't wrapped in a try/except",
 "P2","Medium","Manual","State-dependent (needs a clean environment) - see DEFECT_LOG BUG-07"],

["TC-DASH-01","Dashboard","Positive","Dashboard shows the logged-in user's name",
 "User logged in","1. Log in successfully\n2. Observe dashboard heading","N/A",
 "'Welcome, <name>' is displayed","P2","N/A (positive)","Manual","Depends on TC-LG-01"],

["TC-FILE-01","File Management","Positive","Upload a file while logged in",
 "User logged in, on /dashboard","1. Click 'Upload File'\n2. Choose a file\n3. Click Upload",
 "Any file <server upload limit>","'File uploaded successfully!' flash; file appears under /files",
 "P1","N/A (positive)","Manual","Depends on TC-LG-01"],

["TC-FILE-02","File Management","Negative","Submit upload with no file selected",
 "User logged in, upload modal open","1. Click Upload without choosing a file","N/A",
 "'No file selected.' flash","P3","Low","Manual","Client 'required' attr may block this before reaching server - verify both paths"],

["TC-FILE-03","File Management","Boundary","Upload a very large file",
 "User logged in","1. Attempt to upload a file larger than typical server limits (e.g. 100MB+)",
 "Large file","Expected: graceful rejection with a clear error. Actual: UNVERIFIED - app.py sets no MAX_CONTENT_LENGTH, " 
 "so Flask's default (unlimited) behaviour applies; large uploads may exhaust server memory/disk",
 "P2","Medium","Manual","See DEFECT_LOG BUG-10 (no upload size/type validation)"],

["TC-FILE-04","File Management","Security","Upload a file with no extension / executable file / script",
 "User logged in","1. Upload a .exe, .sh, or extension-less file","Malicious file type",
 "Expected: rejected or sanitized. Actual: accepted as-is, no content-type or extension whitelist enforced",
 "P2","Medium","Manual","See DEFECT_LOG BUG-10"],

["TC-FILE-05","File Management","Positive","View uploaded files list",
 "User logged in, has uploaded at least one file","1. Click 'View Files'","N/A",
 "All files under this user's folder are listed with working Download links","P2","N/A (positive)","Manual","Depends on TC-FILE-01"],

["TC-FILE-06","File Management","Positive","Download own file",
 "User logged in, has an uploaded file","1. Click Download next to a file","N/A",
 "Browser downloads the file successfully","P2","N/A (positive)","Manual","Depends on TC-FILE-01"],

["TC-NAV-01","Navigation","Positive","Home page loads with correct heading",
 "App running","1. Navigate to /","N/A","Heading mentions 'Password Authentication using Eye Blink'",
 "P3","N/A (positive)","Automated (UI)","tests.NavigationTests#homePageLoads"],

["TC-NAV-02","Navigation","Positive","'Get Started' link navigates to Signup",
 "On home page","1. Click 'Get Started'","N/A","Navigates to /signup","P3","N/A (positive)",
 "Automated (UI)","tests.NavigationTests#getStartedGoesToSignup"],

["TC-NAV-03","Navigation","Positive","Nav bar Login link navigates to Login",
 "On home page","1. Click 'Login' in nav","N/A","Navigates to /login","P3","N/A (positive)",
 "Automated (UI)","tests.NavigationTests#loginNavGoesToLogin"],

["TC-LOGOUT-01","Logout","Positive","Logout clears the session",
 "User logged in","1. Click Logout\n2. Then navigate directly to /dashboard","N/A",
 "Session is cleared, 'Logged out.' flash shown; subsequent /dashboard request redirects to /login",
 "P1","N/A (positive)","Manual","Depends on TC-LG-01"],
]

ws.append(headers)
for col in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
ws.row_dimensions[1].height = 24

for r in rows:
    ws.append(r)

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
    for cell in row:
        cell.font = BODY_FONT
        cell.alignment = WRAP
        cell.border = BORDER

widths = [12, 15, 14, 34, 26, 40, 24, 46, 8, 14, 20, 40]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"

# Summary sheet
summary = wb.create_sheet("Summary")
summary["A1"] = "Eye-Blink Authentication System - Test Case Summary"
summary["A1"].font = Font(name="Arial", bold=True, size=13, color="1F3864")
summary.merge_cells("A1:D1")

summary_rows = [
    ("Total test cases", len(rows)),
    ("Automated (UI - Selenium/TestNG)", sum(1 for r in rows if "Automated (UI)" in r[10] or "Automated (UI + API)" in r[10])),
    ("Automated (API - REST Assured)", sum(1 for r in rows if "Automated (API)" in r[10] or "Automated (UI + API)" in r[10])),
    ("Manual", sum(1 for r in rows if r[10] == "Manual")),
    ("Security-focused test cases", sum(1 for r in rows if "Security" in r[2])),
    ("Critical severity findings", sum(1 for r in rows if r[9] == "Critical")),
]
for i, (label, val) in enumerate(summary_rows, start=3):
    summary.cell(row=i, column=1, value=label).font = Font(name="Arial", bold=True, size=10)
    summary.cell(row=i, column=2, value=val).font = BODY_FONT

summary.column_dimensions["A"].width = 38
summary.column_dimensions["B"].width = 12

note = summary.cell(row=len(summary_rows) + 5, column=1,
    value="See DEFECT_LOG.xlsx for the full list of defects found during test design, "
          "including a Critical-severity IDOR on /download (BUG-03).")
note.font = Font(name="Arial", italic=True, size=9, color="666666")
note.alignment = Alignment(wrap_text=True)
summary.merge_cells(start_row=len(summary_rows) + 5, start_column=1, end_row=len(summary_rows) + 5, end_column=6)

wb.move_sheet("Summary", offset=-1)
wb.save("TEST_CASES.xlsx")
print("saved")
